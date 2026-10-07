"""
One-time authoring script that produces the .docx templates referenced by
corpus/templates/manifest.json. Each {{ field }} placeholder is written as a
single run so docxtpl (Jinja2-over-docx) can substitute it cleanly at render
time.

Run from the backend/ directory (re-run after editing any template below):
    python scripts/generate_sample_templates.py

Real templates supplied later should follow the same convention: plain
paragraphs with `{{ field_key }}` placeholders matching manifest.json.
"""
import os

from docx import Document
from docx.enum.text import WD_ALIGN_PARAGRAPH
from docx.shared import Pt

_OUT_DIR = os.path.join(os.path.dirname(__file__), "..", "corpus", "templates")


def _new_doc():
    doc = Document()
    style = doc.styles["Normal"]
    style.font.name = "Times New Roman"
    style.font.size = Pt(12)
    return doc


def _title(doc, text):
    p = doc.add_heading(text, level=0)
    p.alignment = WD_ALIGN_PARAGRAPH.CENTER


def _centered(doc, text, bold=False):
    p = doc.add_paragraph()
    p.alignment = WD_ALIGN_PARAGRAPH.CENTER
    p.add_run(text).bold = bold
    return p


def _signature(doc, label):
    doc.add_paragraph()
    sig = doc.add_paragraph("_________________________")
    sig.alignment = WD_ALIGN_PARAGRAPH.RIGHT
    lbl = doc.add_paragraph(label)
    lbl.alignment = WD_ALIGN_PARAGRAPH.RIGHT


def build_bail_application():
    doc = _new_doc()
    _title(doc, "APPLICATION FOR BAIL")

    _centered(doc, "IN THE {{ court_name }}", bold=True)
    doc.add_paragraph("Case No.: {{ case_no }}    FIR No.: {{ fir_no }}")
    doc.add_paragraph("Police Station: {{ police_station }}")
    doc.add_paragraph("Sections Charged: {{ sections_charged }}")
    doc.add_paragraph()

    doc.add_paragraph("BETWEEN:")
    doc.add_paragraph("{{ applicant_name }}, residing at {{ applicant_address }} ... Applicant")
    doc.add_paragraph("AND")
    doc.add_paragraph("{{ filed_against }} ... Respondent")
    doc.add_paragraph()

    doc.add_heading(
        "APPLICATION UNDER SECTION 483 OF THE BHARATIYA NAGARIK SURAKSHA SANHITA, 2023 "
        "(CORRESPONDING TO SECTION 439 CrPC) FOR GRANT OF BAIL",
        level=1,
    )
    doc.add_paragraph("The Applicant above named respectfully submits as follows:")
    doc.add_paragraph("1. That the Applicant was arrested on {{ date_of_arrest }} in connection with the above case.")
    doc.add_paragraph("2. Grounds for bail:")
    doc.add_paragraph("{{ grounds_for_bail }}")
    doc.add_paragraph(
        "3. That the Applicant undertakes to abide by all conditions the Hon'ble Court may impose "
        "and shall not tamper with evidence or influence witnesses."
    )
    doc.add_paragraph(
        "PRAYER: It is therefore most respectfully prayed that this Hon'ble Court may be pleased to "
        "release the Applicant on bail in the interest of justice."
    )
    doc.add_paragraph()

    doc.add_paragraph("Place: {{ filing_place }}")
    doc.add_paragraph("Date: {{ filing_date }}")
    _signature(doc, "(Applicant / Counsel for Applicant)")

    doc.save(os.path.join(_OUT_DIR, "bail_application.docx"))


def build_affidavit():
    doc = _new_doc()
    _title(doc, "AFFIDAVIT")

    doc.add_paragraph(
        "I, {{ deponent_name }}, {{ deponent_relation }}, aged about {{ deponent_age }} years, "
        "residing at {{ deponent_address }}, do hereby solemnly affirm and state as follows:"
    )
    doc.add_paragraph()
    doc.add_paragraph("{{ statement_content }}")
    doc.add_paragraph()
    doc.add_paragraph(
        "I state that the contents of this affidavit are true and correct to the best of my "
        "knowledge and belief, and nothing material has been concealed therefrom."
    )
    doc.add_paragraph()

    doc.add_paragraph("Place: {{ place }}")
    doc.add_paragraph("Date: {{ date }}")
    _signature(doc, "(Deponent)")

    doc.add_paragraph()
    doc.add_heading("VERIFICATION", level=1)
    doc.add_paragraph(
        "Verified at {{ place }} on {{ date }} that the contents of the above affidavit are true "
        "and correct to my knowledge and belief and nothing has been concealed therefrom."
    )
    _signature(doc, "(Deponent)")

    doc.save(os.path.join(_OUT_DIR, "affidavit.docx"))


def build_deposit_legal_notice():
    doc = _new_doc()
    _title(doc, "LEGAL NOTICE")
    _centered(doc, "(Demand for Refund of Security Deposit)", bold=True)
    doc.add_paragraph()

    doc.add_paragraph("Date: {{ notice_date }}")
    doc.add_paragraph()
    doc.add_paragraph("From:")
    doc.add_paragraph("{{ tenant_name }}")
    doc.add_paragraph("{{ tenant_address }}")
    doc.add_paragraph()
    doc.add_paragraph("To:")
    doc.add_paragraph("{{ landlord_name }}")
    doc.add_paragraph("{{ landlord_address }}")
    doc.add_paragraph()
    p = doc.add_paragraph()
    p.add_run("Subject: ").bold = True
    p.add_run("Demand for refund of security deposit of Rs. {{ deposit_amount }} for the premises at {{ premises_address }}")
    doc.add_paragraph()

    doc.add_paragraph("Sir/Madam,")
    doc.add_paragraph(
        "1. I was your tenant in the premises at {{ premises_address }} under a rent agreement dated "
        "{{ agreement_date }}, and paid you a security deposit of Rs. {{ deposit_amount }} at the start of the tenancy."
    )
    doc.add_paragraph(
        "2. I vacated the premises and handed over vacant possession to you on {{ vacated_date }}, "
        "after clearing all rent and charges due up to that date."
    )
    doc.add_paragraph("3. Facts regarding the withholding of the deposit:")
    doc.add_paragraph("{{ dispute_details }}")
    doc.add_paragraph(
        "4. Under Section 108(m) of the Transfer of Property Act, 1882, a tenant is bound to return the "
        "property in the condition in which it was received, subject only to changes caused by reasonable "
        "wear and tear. Ordinary wear and tear from normal use is therefore not a ground to withhold the "
        "deposit, and any deduction must relate to an actual liability of the tenant supported by evidence."
    )
    doc.add_paragraph(
        "5. I hereby call upon you to refund the security deposit of Rs. {{ deposit_amount }} within "
        "{{ response_days }} days of receipt of this notice, by bank transfer or demand draft in my name."
    )
    doc.add_paragraph(
        "6. If you fail to do so, I shall be constrained to initiate appropriate legal proceedings against "
        "you for recovery of the said amount along with interest and costs, entirely at your risk as to "
        "costs and consequences."
    )
    doc.add_paragraph()
    doc.add_paragraph("A copy of this notice has been retained for further action.")
    _signature(doc, "({{ tenant_name }})")

    doc.save(os.path.join(_OUT_DIR, "deposit_legal_notice.docx"))


def build_rti_application():
    doc = _new_doc()
    _title(doc, "APPLICATION UNDER THE RIGHT TO INFORMATION ACT, 2005")
    _centered(doc, "(Under Section 6(1) of the RTI Act, 2005)")
    doc.add_paragraph()

    doc.add_paragraph("To,")
    doc.add_paragraph("The Public Information Officer,")
    doc.add_paragraph("{{ public_authority }}")
    doc.add_paragraph("{{ authority_address }}")
    doc.add_paragraph()

    doc.add_paragraph("1. Full name of the applicant: {{ applicant_name }}")
    doc.add_paragraph("2. Address for correspondence: {{ applicant_address }}")
    doc.add_paragraph("3. Phone / Email (optional): {{ applicant_contact }}")
    doc.add_paragraph("4. Particulars of the information sought:")
    doc.add_paragraph("{{ information_sought }}")
    doc.add_paragraph("5. Period to which the information relates: {{ period }}")
    doc.add_paragraph("6. Application fee: {{ fee_details }}")
    doc.add_paragraph("7. Whether the applicant is below the poverty line: {{ bpl_status }}")
    doc.add_paragraph()
    doc.add_paragraph(
        "I state that the information sought does not fall within the restrictions contained in "
        "Sections 8 and 9 of the RTI Act and, to the best of my knowledge, pertains to your office. "
        "I request that the information be provided within 30 days as required under Section 7 of the Act."
    )
    doc.add_paragraph()
    doc.add_paragraph("Place: {{ place }}")
    doc.add_paragraph("Date: {{ date }}")
    _signature(doc, "({{ applicant_name }})")

    doc.save(os.path.join(_OUT_DIR, "rti_application.docx"))


def build_consumer_complaint():
    doc = _new_doc()
    _title(doc, "CONSUMER COMPLAINT")
    _centered(doc, "BEFORE THE DISTRICT CONSUMER DISPUTES REDRESSAL COMMISSION, {{ district }}", bold=True)
    _centered(doc, "(Complaint under Section 35 of the Consumer Protection Act, 2019)")
    doc.add_paragraph()

    doc.add_paragraph("Consumer Complaint No. ______ of 20___")
    doc.add_paragraph()
    doc.add_paragraph("{{ complainant_name }}, residing at {{ complainant_address }} ... Complainant")
    doc.add_paragraph("VERSUS")
    doc.add_paragraph("{{ opposite_party_name }}, having its office at {{ opposite_party_address }} ... Opposite Party")
    doc.add_paragraph()

    doc.add_paragraph("The Complainant respectfully submits as follows:")
    doc.add_paragraph(
        "1. That the Complainant purchased / availed {{ product_or_service }} from the Opposite Party on "
        "{{ purchase_date }} for a consideration of Rs. {{ amount_paid }}, and is a 'consumer' within the "
        "meaning of Section 2(7) of the Consumer Protection Act, 2019."
    )
    doc.add_paragraph("2. Facts of the complaint (defect / deficiency in service / unfair trade practice):")
    doc.add_paragraph("{{ complaint_facts }}")
    doc.add_paragraph(
        "3. That the cause of action arose on {{ purchase_date }} and thereafter, within the jurisdiction "
        "of this Hon'ble Commission, and the complaint is filed within the limitation period of two years."
    )
    doc.add_paragraph("4. Relief sought:")
    doc.add_paragraph("{{ relief_sought }}")
    doc.add_paragraph()
    doc.add_paragraph(
        "PRAYER: It is therefore prayed that this Hon'ble Commission may be pleased to direct the "
        "Opposite Party to grant the relief sought above, together with compensation and costs of "
        "this complaint, in the interest of justice."
    )
    doc.add_paragraph()
    doc.add_paragraph("Place: {{ place }}")
    doc.add_paragraph("Date: {{ date }}")
    _signature(doc, "(Complainant)")

    doc.add_paragraph()
    doc.add_heading("VERIFICATION", level=1)
    doc.add_paragraph(
        "I, {{ complainant_name }}, the Complainant above named, do hereby verify that the contents of "
        "this complaint are true and correct to my knowledge and belief and nothing material has been concealed."
    )
    _signature(doc, "(Complainant)")

    doc.save(os.path.join(_OUT_DIR, "consumer_complaint.docx"))


if __name__ == "__main__":
    os.makedirs(_OUT_DIR, exist_ok=True)
    build_bail_application()
    build_affidavit()
    build_deposit_legal_notice()
    build_rti_application()
    build_consumer_complaint()
    print(f"Wrote 5 templates to {os.path.abspath(_OUT_DIR)}")
