"""
Authoring script for the consumer-case template pack (corpus/templates/consumer_*.docx).

Each template mirrors the structure and wording of the reference formats the
team supplied (Consumer-Templates pack, numbered 1-7), with two deliberate
changes:
  - statutory references updated from the repealed Consumer Protection Act,
    1986 to the Consumer Protection Act, 2019 (complaint s.12 -> s.35,
    execution s.25/27 -> s.71/72, "Forum" -> "Commission");
  - the two "stand-in" formats (#1 cheque-dishonour notice, #6 Income Tax
    Tribunal condonation affidavit) are adapted to consumer disputes,
    keeping their structure.
#4 (Rejoinder) was not in the pack and is written in the same style.

Blanks in the reference formats become {{ field }} placeholders. Blanks the
court registry fills in (complaint / application numbers) fall back to
underscores when left empty.

Run from the backend/ directory:
    python scripts/generate_consumer_templates.py
"""
import os

from docx import Document
from docx.enum.text import WD_ALIGN_PARAGRAPH, WD_BREAK
from docx.shared import Pt

_OUT_DIR = os.path.join(os.path.dirname(__file__), "..", "corpus", "templates")

J = WD_ALIGN_PARAGRAPH.JUSTIFY
C = WD_ALIGN_PARAGRAPH.CENTER
R = WD_ALIGN_PARAGRAPH.RIGHT

BLANK = "'______'"


def _new_doc():
    doc = Document()
    style = doc.styles["Normal"]
    style.font.name = "Times New Roman"
    style.font.size = Pt(12)
    style.paragraph_format.space_after = Pt(6)
    return doc


def p(doc, text="", align=J, bold=False, underline=False):
    para = doc.add_paragraph()
    para.alignment = align
    if text:
        run = para.add_run(text)
        run.bold = bold
        run.underline = underline
    return para


def page_break(doc):
    doc.add_paragraph().add_run().add_break(WD_BREAK.PAGE)


def cause_title(doc, commission_line, number_line, left_label, right_label, left_key, right_key):
    """'BEFORE THE ... / No. / X ... Complainant / Versus / Y ... Respondents' block used by every court format."""
    p(doc, commission_line, C, bold=True)
    p(doc, number_line, C)
    p(doc)
    p(doc, "{{ " + left_key + " }} ..... " + left_label)
    p(doc, "Versus", C)
    p(doc, "{{ " + right_key + " }} ..... " + right_label)
    p(doc)


def signature(doc, label, through_advocate=True):
    # Keep the two lines before it (Place / Date) and the whole block on one
    # page, so a signature never lands alone on a new page.
    block = doc.paragraphs[-2:]
    block += [p(doc), p(doc, "......................................", R), p(doc, label, R)]
    if through_advocate:
        block.append(p(doc, "Through, {{ advocate_name or '______' }}, Advocate", R))
    for para in block[:-1]:
        para.paragraph_format.keep_with_next = True


# ---------------------------------------------------------------------------
# 1. Legal notice (adapted from the NI Act s.138 cheque-dishonour stand-in)
# ---------------------------------------------------------------------------

def build_legal_notice():
    doc = _new_doc()
    p(doc, "LEGAL NOTICE FOR DEFICIENCY IN SERVICE / DEFECT IN GOODS UNDER THE CONSUMER PROTECTION ACT, 2019", C, bold=True)
    p(doc, "REGISTERED A.D.", C, bold=True)
    p(doc, "To,")
    p(doc, "{{ opposite_party_name }}")
    p(doc, "{{ opposite_party_address }}")
    p(doc)
    p(doc, "Sub:- Notice for {{ grievance_type }} in respect of {{ product_or_service }} purchased/availed by our client.", bold=True)
    p(doc, "Dear Sir,")
    p(doc,
      "Under instructions and authority from our client {{ client_name }} residing at {{ client_address }}, "
      "we serve upon you the following notice of demand under the Consumer Protection Act, 2019.")
    p(doc,
      "1. That our client purchased/availed from you {{ product_or_service }} vide invoice/receipt bearing "
      "no. {{ invoice_no or '______' }} dated {{ purchase_date }} for Rs. {{ amount_paid }}.")
    p(doc, "2. That the said goods/service suffered from the following defect/deficiency:")
    p(doc, "{{ defect_details }}")
    p(doc,
      "3. That this fact was brought to your notice by our client vide complaint/letter dated "
      "{{ first_complaint_date }}.")
    p(doc, "4. That in reply to our client's complaint, your response was as under:")
    p(doc, "{{ opposite_party_response or 'No response was received from you.' }}")
    p(doc,
      "5. That the defect/deficiency has not been remedied by you till date, and our client continues to "
      "suffer loss on account of the same.")
    p(doc,
      "6. That thereafter in spite of many telephonic reminders and personal visits by our client, you have "
      "failed to resolve the grievance of our client.")
    p(doc,
      "7. That on account of the above facts, your acts and omissions amount to defect in goods, deficiency "
      "in service and/or unfair trade practice within the meaning of Sections 2(10), 2(11) and 2(47) of the "
      "Consumer Protection Act, 2019, for which you are liable to compensate our client.")
    p(doc,
      "8. Under the circumstances, we call upon you to {{ demand }} along with interest @ "
      "{{ interest_rate or '______' }} % per annum within a period of {{ notice_period or '15 (fifteen)' }} days "
      "from the date of receipt of this notice, failing which we will be bound to file a complaint against you "
      "under Section 35 of the Consumer Protection Act, 2019 before the competent Consumer Disputes "
      "Redressal Commission at your risk and cost.")
    p(doc,
      "This is without prejudice to all other legal rights and remedies available to our client for the "
      "above-stated purpose.")
    p(doc, "Kindly take notice.")
    p(doc)
    p(doc, "Date: {{ notice_date }}")
    p(doc, "Yours faithfully,", R)
    p(doc)
    p(doc, "({{ advocate_name or '______' }})", R)
    p(doc, "Advocate for {{ client_name }}", R)
    doc.save(os.path.join(_OUT_DIR, "consumer_legal_notice.docx"))


# ---------------------------------------------------------------------------
# 2. Consumer complaint
# ---------------------------------------------------------------------------

def build_complaint():
    doc = _new_doc()
    cause_title(
        doc,
        "BEFORE THE DISTRICT CONSUMER DISPUTES REDRESSAL COMMISSION AT {{ commission_place }}",
        "Complaint No.: {{ complaint_no or '______' }}",
        "Complainant", "Opposite Parties",
        "complainant_details", "opposite_party_details",
    )
    p(doc, "Complaint under Section 35 of the Consumer Protection Act, 2019", C, bold=True)
    p(doc, "Respectfully Sheweth:")
    p(doc,
      "1. That the relevant facts and the facts leading to filing of present complaint are given in brief "
      "hereunder in chronological order for the convenience of your honor :-")
    p(doc, "{{ events }}")
    p(doc, "GROUNDS", C, bold=True)
    p(doc,
      "2. That the above being the background of the case, the complainant is entitled to maintain and file the "
      "present complaint before this Hon'ble Commission on the following grounds amongst others, each one of "
      "which is without prejudice to others, and seeks the indulgence of this Hon'ble Commission to issue "
      "directions or orders to the opposite parties deemed appropriate :-")
    p(doc,
      "(a) That such an act and omission on the part of the opposite parties amounts to deficiency in service, "
      "restrictive trade practice and unfair trade practice under the purview of the Consumer Protection Act, 2019.")
    p(doc,
      "(b) That the complainant has suffered lot of avoidable mental agonies, physical harassment and financial "
      "loss, irreparable loss and injury at the hands of the opposite parties for which the complainant deserves "
      "to be amply and suitably compensated in terms of money in the interest of justice.")
    p(doc,
      "(c) That the opposite parties deserve to be saddled with heavy and extra-ordinary costs so that the "
      "opposite parties can be deterred from adopting such mal-practice in future.")
    p(doc,
      "(d) That the complainant herein deserves to be amply compensated in terms of money for the inconvenience, "
      "traveling expenses, for mental agonies, damages, delayed service/no service, interest, and loss suffered "
      "due to negligence and dereliction of duties by the opposite parties.")
    p(doc, "3. That the complaint is within the period of limitation.")
    p(doc,
      "4. That this Hon'ble Commission has jurisdiction to entertain and to try the complaint as "
      "{{ jurisdiction_reason }}, within the jurisdiction of this Hon'ble Commission, and the value of the goods "
      "or services paid as consideration is within the pecuniary limits of this Hon'ble Commission. It is "
      "submitted that the present complaint is being filed without prejudice to other rights or remedies "
      "available to the complainant.")
    p(doc,
      "5. It is, therefore, prayed that this Hon'ble Commission may be pleased to pass the following orders, "
      "directions and grant the following reliefs in favour of the complainant in the interest of justice :-")
    p(doc,
      "(a) Direct the opposite parties to {{ main_relief }} along with interest @ {{ interest_rate or '______' }} % "
      "from the date of payment till realisation;")
    p(doc,
      "(b) Saddle the opposite parties with special and extra-ordinary costs as deemed fit so as to deter them "
      "from adopting such malpractice in future;")
    p(doc,
      "(c) Award a compensation of Rs. {{ compensation_amount }} to the complainant on account of mental agony, "
      "physical harassment, and financial loss;")
    p(doc, "(d) Allow the cost of this complaint;")
    p(doc,
      "(e) Pass such other orders in favour of the complainant as deemed fit and proper in the facts and "
      "circumstances of the case.")
    p(doc, "AND FOR THIS ACT OF KINDNESS, THE HUMBLE COMPLAINANT AS IN DUTY BOUND, SHALL EVER PRAY", C, bold=True)
    p(doc, "Place: {{ place }}")
    p(doc, "Date: {{ date }}")
    signature(doc, "Complainant")
    doc.save(os.path.join(_OUT_DIR, "consumer_complaint.docx"))


# ---------------------------------------------------------------------------
# 3. Affidavit in support of the complaint
# ---------------------------------------------------------------------------

def build_complaint_affidavit():
    doc = _new_doc()
    cause_title(
        doc,
        "BEFORE THE DISTRICT CONSUMER DISPUTES REDRESSAL COMMISSION AT {{ commission_place }}",
        "Complaint No.: {{ complaint_no or '______' }}",
        "Complainant", "Opposite Parties",
        "complainant_name", "opposite_party_name",
    )
    p(doc, "Affidavit in support of complaint under Section 35 of the Consumer Protection Act, 2019", C, bold=True)
    p(doc,
      "I, {{ deponent_name }}, {{ deponent_relation }}, aged {{ deponent_age }} years, residing at "
      "{{ deponent_address }}, do hereby solemnly affirm and declare as under:-")
    p(doc, "1. That the accompanying complaint has been prepared under my instructions.")
    p(doc,
      "2. That the contents of paras 1 to {{ last_para or '5' }} of the complaint are correct and true to the "
      "best of my knowledge.")
    p(doc,
      "3. That I further solemnly affirm and declare that this affidavit of mine is correct and true to the best "
      "of my knowledge and no part of it is false and nothing material has been concealed therefrom.")
    p(doc)
    p(doc, "Affirmed at {{ place }} this the {{ date }}.")
    p(doc)
    p(doc, "Deponent", R, bold=True)
    doc.save(os.path.join(_OUT_DIR, "consumer_complaint_affidavit.docx"))


# ---------------------------------------------------------------------------
# 4. Rejoinder (not in the pack; written in the same style)
# ---------------------------------------------------------------------------

def build_rejoinder():
    doc = _new_doc()
    cause_title(
        doc,
        "BEFORE THE DISTRICT CONSUMER DISPUTES REDRESSAL COMMISSION AT {{ commission_place }}",
        "Complaint No.: {{ complaint_no or '______' }}",
        "Complainant", "Opposite Parties",
        "complainant_name", "opposite_party_name",
    )
    p(doc,
      "Rejoinder on behalf of the Complainant to the Written Version filed by the Opposite Party "
      "dated {{ written_version_date }}", C, bold=True)
    p(doc, "Respectfully Sheweth:")
    p(doc,
      "1. That the contents of the written version filed by the opposite party are wrong and denied, except "
      "what is specifically admitted hereinafter. Nothing contained in the written version shall be deemed to "
      "be admitted by the complainant merely for want of specific denial.")
    p(doc,
      "2. That the contents of the complaint are reiterated and reaffirmed as true and correct, and the same "
      "may kindly be read as part and parcel of this rejoinder.")
    p(doc, "3. That the reply of the complainant to the written version is as under :-")
    p(doc, "{{ parawise_reply }}")
    p(doc,
      "4. That the opposite party has failed to place on record any document or material to rebut the "
      "averments of the complaint, and the defence taken in the written version is an afterthought, taken "
      "only to escape liability for the deficiency in service and unfair trade practice complained of.")
    p(doc,
      "5. It is, therefore, prayed that this Hon'ble Commission may be pleased to allow the complaint as prayed "
      "for, with costs, in the interest of justice.")
    p(doc, "AND FOR THIS ACT OF KINDNESS, THE HUMBLE COMPLAINANT AS IN DUTY BOUND, SHALL EVER PRAY", C, bold=True)
    p(doc, "Place: {{ place }}")
    p(doc, "Date: {{ date }}")
    signature(doc, "Complainant")
    p(doc)
    p(doc, "VERIFICATION", C, bold=True)
    p(doc,
      "Verified at {{ place }} on {{ date }} that the contents of paras 1 to 5 of the above rejoinder are true "
      "and correct to my knowledge and belief, no part of it is false and nothing material has been concealed "
      "therein.")
    p(doc)
    p(doc, "Complainant", R, bold=True)
    doc.save(os.path.join(_OUT_DIR, "consumer_rejoinder.docx"))


# ---------------------------------------------------------------------------
# 5. Evidence by way of affidavit
# ---------------------------------------------------------------------------

def build_evidence_affidavit():
    doc = _new_doc()
    cause_title(
        doc,
        "BEFORE THE HON'BLE {{ commission_name }} AT {{ commission_place }}",
        "Complaint No.: {{ complaint_no or '______' }}",
        "Complainant", "Opposite Parties",
        "complainant_name", "opposite_party_name",
    )
    p(doc, "Evidence by the complainant by way of Affidavit in support of the complaint", C, bold=True)
    p(doc,
      "I, {{ deponent_name }}, aged {{ deponent_age }} years, presently {{ deponent_occupation }}, "
      "do hereby solemnly affirm and declare as under:-")
    p(doc,
      "1. That I have filed documents with the original complaint which may kindly be read as evidence. "
      "The contents of the original complaint may kindly also be read as part of evidence in the above complaint.")
    p(doc,
      "2. That I further solemnly affirm and declare that the contents of this affidavit of mine are correct and "
      "true, no part of it is false and nothing material has been concealed therein.")
    p(doc)
    p(doc, "Affirmed here at {{ place }} this the {{ date }}.")
    p(doc)
    p(doc, "Deponent", R, bold=True)
    doc.save(os.path.join(_OUT_DIR, "consumer_evidence_affidavit.docx"))


# ---------------------------------------------------------------------------
# 6. Condonation of delay (adapted from the Income Tax Tribunal stand-in)
# ---------------------------------------------------------------------------

def build_condonation_affidavit():
    doc = _new_doc()
    p(doc, "AFFIDAVIT FOR CONDONATION OF DELAY IN FILING {{ filing_type | upper }} BEFORE THE "
           "{{ commission_name | upper }}", C, bold=True)
    p(doc, "Before the {{ commission_name }}")
    p(doc, "at {{ commission_place }}")
    p(doc, "In the matter of {{ matter_title }}")
    p(doc, "Affidavit of Mr./Ms. {{ deponent_name }}")
    p(doc, "Aged {{ deponent_age }} years,")
    p(doc, "{{ deponent_capacity }}.")
    p(doc, "That I the above named deponent, am well conversant with the facts deposed to below.")
    p(doc,
      "1. That the cause of action for the present {{ filing_type }} arose on {{ cause_of_action_date }}, "
      "{{ cause_of_action_details }}.")
    p(doc,
      "2. That the time for filing of the {{ filing_type }} was to expire on {{ limitation_expiry_date }}, "
      "being the period of limitation prescribed under {{ limitation_provision }} of the Consumer Protection "
      "Act, 2019.")
    p(doc, "3. That the {{ filing_type }} could not be filed within time for the following reasons:")
    p(doc, "{{ delay_reasons }}")
    p(doc,
      "4. That the {{ filing_type }} has been filed on {{ filing_date }} in the office of the "
      "{{ commission_name }}, and there is a delay of {{ delay_days }} days in filing the same.")
    p(doc,
      "5. That the said delay is neither intentional nor deliberate and was caused for the reasons stated "
      "above, which constitute sufficient cause for not filing the {{ filing_type }} within the prescribed "
      "period, and the deponent prays that the delay may kindly be condoned.")
    p(doc, "DEPONENT", R, bold=True)
    p(doc, "Signed at {{ place }}")
    p(doc, "this {{ date }}.")
    p(doc)
    p(doc, "VERIFICATION", C, bold=True)
    p(doc,
      "I, {{ deponent_name }} the above named deponent do hereby verify on oath that the contents of the "
      "affidavit above are true to my personal knowledge and nothing material has been concealed or falsely "
      "stated. Verified at {{ place }} this {{ date }}.")
    p(doc)
    p(doc, "DEPONENT", R, bold=True)
    doc.save(os.path.join(_OUT_DIR, "consumer_condonation_affidavit.docx"))


# ---------------------------------------------------------------------------
# 7. Execution application (+ supporting affidavit, as in the pack)
# ---------------------------------------------------------------------------

def build_execution_application():
    doc = _new_doc()
    number_line = "Application No.: {{ application_no or '______' }} in Complaint No.: {{ complaint_no or '______' }}"
    cause_title(
        doc,
        "BEFORE THE HON'BLE DISTRICT CONSUMER DISPUTES REDRESSAL COMMISSION AT {{ commission_place }}",
        number_line,
        "Applicant/Complainant", "Respondents/Non-applicants",
        "applicant_name", "respondent_name",
    )
    p(doc, "Application under Sections 71 and 72 of the Consumer Protection Act, 2019", C, bold=True)
    p(doc, "Respectfully Sheweth:")
    p(doc,
      "1. That the applicant hereinabove has filed a complaint No. {{ complaint_no or '______' }}, which your honour were "
      "pleased to decide on {{ order_date }} [Annexure C-1] with the following directions to the respondents :-")
    p(doc, "\"{{ order_directions }}\"")
    p(doc,
      "2. That it is apparent from the above orders passed by this Hon'ble Commission that the order was "
      "comprehensive, unambiguous and clear in directing the respondents to comply with the same within "
      "{{ compliance_days }} days.")
    p(doc,
      "3. That the said orders of this Hon'ble Commission were duly served on the respondents "
      "{{ service_details }}. Legally the service of the above orders is complete on the respondents.")
    p(doc,
      "4. That despite the orders of this Hon'ble Commission having duly been served on and brought to the notice "
      "of the respondents, the respondents have not complied with the same within the specified period.")
    p(doc,
      "5. That the respondents have thus conducted themselves in the most reprehensible manner showing scant "
      "regard to the orders of this Hon'ble Commission which conduct on their part is palpably, manifestly and "
      "gravely contumacious and makes the contemnors/respondents liable to be dealt with in accordance with the "
      "law and punished severely for willfully, deliberately and intentionally flouting, defying and disobeying "
      "the authority of this Hon'ble Commission.")
    p(doc,
      "6. That it is apparent from the material placed on record that no iota of doubt is left regarding the "
      "contemnors having scant regard and utter disregard to the sacrosanct orders of this Hon'ble Commission, "
      "and the contemnors are prima facie guilty of having deliberately and willfully disobeyed the orders of "
      "this Hon'ble Commission and they deserve to be dealt with and punished severely in accordance with "
      "Section 72 of the Consumer Protection Act, 2019, as by their such unbecoming act they have tended to "
      "lower the authority of this Hon'ble Commission and tended to interfere with the due process of the "
      "judicial proceedings and administration of justice as their this conduct may make execution of order "
      "infructuous. Their such conduct is manifest of their malafide intentions throughout.")
    p(doc,
      "7. That the interest of justice demands that the contemnors/respondents are dealt with and punished "
      "severely in accordance with the law for their willful, deliberate and intentional disobedience of the "
      "orders of this Hon'ble Commission. Otherwise, if the contemnors are let go scot free the litigants will "
      "lose their confidence and faith in the judicial proceedings.")
    p(doc,
      "8. It is, therefore, most respectfully prayed that this Hon'ble Commission may be pleased to grant the "
      "following reliefs/pass the following orders :-")
    p(doc,
      "(a) That the contemnors/respondents be severely dealt with and punished in accordance with Section 72 of "
      "the Consumer Protection Act, 2019 for having willfully, deliberately and intentionally not complied with "
      "the orders of this Hon'ble Commission, thereby lowering the dignity of this Hon'ble Commission;")
    p(doc,
      "(b) The execution proceedings be initiated under Section 71 of the Consumer Protection Act, 2019 for "
      "execution of the above orders of this Hon'ble Commission;")
    p(doc, "(c) The respondents/contemnors be burdened with extra heavy costs;")
    p(doc,
      "(d) Such other orders may also be passed in favour of the complainant as deemed fit and proper by this "
      "Hon'ble Commission in the facts and circumstances of the case.")
    p(doc, "AND FOR THIS ACT OF KINDNESS, THE HUMBLE APPLICANT AS IN DUTY BOUND, SHALL EVER PRAY", C, bold=True)
    p(doc, "Place: {{ place }}")
    p(doc, "Date: {{ date }}")
    signature(doc, "Applicant")

    page_break(doc)
    cause_title(
        doc,
        "BEFORE THE HON'BLE DISTRICT CONSUMER DISPUTES REDRESSAL COMMISSION AT {{ commission_place }}",
        number_line,
        "Applicant/Complainant", "Respondents/Non-applicants",
        "applicant_name", "respondent_name",
    )
    p(doc, "Affidavit in support of Application under Sections 71 and 72 of the Consumer Protection Act, 2019", C, bold=True)
    p(doc, "I, {{ applicant_name }}, do hereby solemnly affirm and declare as under :-")
    p(doc, "1. That the accompanying application has been prepared under my instructions.")
    p(doc,
      "2. That the contents of paras 1 to 8 of the accompanying application are true and correct to the best of "
      "my knowledge.")
    p(doc,
      "3. That I further solemnly affirm and declare that this affidavit of mine is correct and true, no part of "
      "it is false and nothing material has been concealed therein.")
    p(doc)
    p(doc, "Affirmed at {{ place }} this the {{ date }}.")
    p(doc)
    p(doc, "Deponent", R, bold=True)
    doc.save(os.path.join(_OUT_DIR, "consumer_execution_application.docx"))


if __name__ == "__main__":
    os.makedirs(_OUT_DIR, exist_ok=True)
    build_legal_notice()
    build_complaint()
    build_complaint_affidavit()
    build_rejoinder()
    build_evidence_affidavit()
    build_condonation_affidavit()
    build_execution_application()
    print(f"Wrote 7 consumer templates to {os.path.abspath(_OUT_DIR)}")
